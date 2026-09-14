from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


class ExcelExporter:
    def __init__(self, findings, aws, output):
        self.findings = findings
        self.aws = aws
        self.output = output

    def export(self):
        workbook = Workbook()

        summary = workbook.active
        summary.title = "Summary"

        self.write_summary(summary)

        grouped = {}

        for finding in self.findings:
            service = finding["service"]
            grouped.setdefault(service, []).append(finding)

        for service, findings in grouped.items():
            sheet_name = service.replace(" ", "")[:31]
            worksheet = workbook.create_sheet(sheet_name)
            self.write_service_sheet(worksheet, findings)

        findings_sheet = workbook.create_sheet("Findings")
        self.write_findings_sheet(findings_sheet)

        workbook.save(self.output)

    def write_summary(self, worksheet):
        worksheet.append(["AWS Unused Resources Report"])
        worksheet.append([])
        worksheet.append(["Account ID", self.aws.account_id()])
        worksheet.append(["Region", self.aws.region])
        worksheet.append([
            "Generated At",
            datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        ])
        worksheet.append([])
        worksheet.append(["Service", "Findings"])

        services = {}

        for finding in self.findings:
            service = finding["service"]
            services[service] = services.get(service, 0) + 1

        for service, count in sorted(services.items()):
            worksheet.append([service, count])

        worksheet.append([])
        worksheet.append(["Total Findings", len(self.findings)])

        self.format_sheet(worksheet)

    def write_service_sheet(self, worksheet, findings):
        worksheet.append([
            "Resource ID",
            "Resource Name",
            "Finding",
            "Risk",
            "Region",
            "Details"
        ])

        for finding in findings:
            worksheet.append([
                finding["resource_id"],
                finding["resource_name"],
                finding["finding"],
                finding["risk"],
                finding["region"],
                str(finding["details"])
            ])

        self.format_sheet(worksheet)

    def write_findings_sheet(self, worksheet):
        worksheet.append([
            "Service",
            "Resource ID",
            "Resource Name",
            "Finding",
            "Risk",
            "Region",
            "Details"
        ])

        for finding in self.findings:
            worksheet.append([
                finding["service"],
                finding["resource_id"],
                finding["resource_name"],
                finding["finding"],
                finding["risk"],
                finding["region"],
                str(finding["details"])
            ])

        self.format_sheet(worksheet)

    def format_sheet(self, worksheet):
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        for cell in worksheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="1F4E78"
            )
            cell.font = Font(
                bold=True,
                color="FFFFFF"
            )
            cell.alignment = Alignment(
                horizontal="center"
            )

        for row in worksheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True
                )

                if cell.value == "High":
                    cell.fill = PatternFill(
                        fill_type="solid",
                        fgColor="FFC7CE"
                    )

                elif cell.value == "Medium":
                    cell.fill = PatternFill(
                        fill_type="solid",
                        fgColor="FFEB9C"
                    )

                elif cell.value == "Low":
                    cell.fill = PatternFill(
                        fill_type="solid",
                        fgColor="C6EFCE"
                    )

        for column in worksheet.columns:
            max_length = 0
            column_letter = get_column_letter(
                column[0].column
            )

            for cell in column:
                value = str(cell.value or "")
                max_length = max(
                    max_length,
                    len(value)
                )

            worksheet.column_dimensions[
                column_letter
            ].width = min(max_length + 2, 50)

        worksheet.row_dimensions[1].height = 25